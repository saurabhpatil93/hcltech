from pathlib import Path
import hashlib, chromadb
from sentence_transformers import SentenceTransformer
from app.config import CHROMA_DIR,RETRIEVAL_TOP_K
from app.db import connect
_MODEL=None
def embedder():
 global _MODEL
 if _MODEL is None: _MODEL=SentenceTransformer('sentence-transformers/all-MiniLM-L6-v2')
 return _MODEL
def collection():
 CHROMA_DIR.mkdir(parents=True,exist_ok=True); return chromadb.PersistentClient(path=str(CHROMA_DIR)).get_or_create_collection('retail_policy_chunks',metadata={'hnsw:space':'cosine'})
def extract_text(p):
 if p.suffix.lower()=='.pdf':
  import fitz
  d=fitz.open(p); return '\n'.join(f'[Page {i+1}] {x.get_text()}' for i,x in enumerate(d))
 if p.suffix.lower()=='.docx':
  from docx import Document
  d=Document(p); return '\n'.join(x.text for x in d.paragraphs)+'\n'+'\n'.join(' | '.join(c.text for c in row.cells) for t in d.tables for row in t.rows)
 if p.suffix.lower() in ('.txt','.md'): return p.read_text(encoding='utf-8',errors='ignore')
 raise ValueError('Supported formats: PDF, DOCX, TXT, MD')
def chunk_text(text,size=900,overlap=150):
 text=' '.join(text.split()); out=[]; start=0
 while start<len(text):
  end=min(start+size,len(text)); out.append(text[start:end])
  if end==len(text): break
  start=end-overlap
 return out
def ingest_file(path,meta):
 chunks=chunk_text(extract_text(path))
 if not chunks: raise ValueError('No extractable text; OCR scanned documents first')
 col=collection(); ids=[]; metas=[]
 for i,ch in enumerate(chunks):
  ids.append(hashlib.sha256(f"{meta['doc_id']}:{i}:{ch}".encode()).hexdigest()[:32])
  metas.append({'doc_id':meta['doc_id'],'filename':path.name,'section':f'chunk {i+1}','version':meta.get('version','1.0'),'effective_from':meta['effective_from'],'effective_to':meta.get('effective_to',''),'authority_level':int(meta['authority_level']),'category':meta.get('category','GENERAL'),'scope_tiers':meta.get('scope_tiers','ALL'),'synthetic':meta.get('synthetic','N'),'source_url':meta.get('source_url',''),'supersedes':meta.get('supersedes','')})
 vectors=embedder().encode(chunks,normalize_embeddings=True).tolist()
 existing_ids=col.get(where={'doc_id':meta['doc_id']},include=['metadatas'])['ids']
 col.upsert(ids=ids,documents=chunks,metadatas=metas,embeddings=vectors)
 stale_ids=set(existing_ids)-set(ids)
 if stale_ids: col.delete(ids=list(stale_ids))
 with connect() as c: c.execute('INSERT INTO source_register VALUES(?,?,?,?,?,?,?,?,?,?,?) ON CONFLICT(doc_id) DO UPDATE SET filename=excluded.filename,authority_level=excluded.authority_level,effective_from=excluded.effective_from,effective_to=excluded.effective_to,version=excluded.version,supersedes=excluded.supersedes',(meta['doc_id'],path.name,int(meta['authority_level']),meta.get('category','GENERAL'),meta.get('scope_tiers','ALL'),meta['effective_from'],meta.get('effective_to',''),meta.get('version','1.0'),meta.get('source_url',''),meta.get('synthetic','N'),meta.get('supersedes','')))
 return {'doc_id':meta['doc_id'],'chunks_indexed':len(chunks),'status':'indexed'}
def search(q,asof,category='GENERAL',tier='STANDARD',k=None):
 col=collection()
 if not col.count(): return []
 res=col.query(query_embeddings=embedder().encode([q],normalize_embeddings=True).tolist(),n_results=min(k or RETRIEVAL_TOP_K,col.count()),include=['documents','metadatas','distances']); hits=[]
 for doc,m,d in zip(res['documents'][0],res['metadatas'][0],res['distances'][0]):
  if m.get('effective_from','0001-01-01')>asof or (m.get('effective_to') and m['effective_to']<asof): continue
  if m.get('scope_tiers','ALL') not in ('ALL',tier): continue
  hits.append({'text':doc,'metadata':m,'distance':float(d)})
 return hits
